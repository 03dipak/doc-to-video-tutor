"""Filesystem helpers (atomic json write, progress bar, document loading)."""

from __future__ import annotations

import json
import os
import time
from pathlib import Path


def atomic_json_write(path: Path, data: dict) -> None:
    """Write JSON atomically (tmp file + fsync + rename) so a crashed write
    never leaves a truncated .plan.json / .audit.json behind. Same-directory
    rename keeps the move atomic on POSIX."""
    tmp = path.with_name(f".{path.name}.tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp, path)
class _Progress:
    def __init__(self, stage: str, total: int):
        self.stage = stage
        self.total = total
        self.start = time.monotonic()
        self._tick(0)

    def _tick(self, done: int):
        elapsed = time.monotonic() - self.start
        frac = done / self.total if self.total else 0
        bar_len = 20
        filled = int(bar_len * frac)
        bar = "\u2588" * filled + "\u2591" * (bar_len - filled)
        if done > 0 and self.total > 1:
            per_item = elapsed / done
            eta = int(per_item * (self.total - done))
            eta_str = f"{eta // 60}m {eta % 60:02d}s" if eta else "done"
        else:
            eta_str = "..."
        print(
            f"\r  [{self.stage}] |{bar}| {frac:5.0%}  elapsed {elapsed:.0f}s  ETA {eta_str}   ",
            end="", flush=True,
        )

    def done(self):
        elapsed = time.monotonic() - self.start
        print(f"\r  [{self.stage}] \u2713  done in {elapsed:.0f}s{' ' * 40}", flush=True)
def _wait_before_retry(failed_samples: int) -> None:
    """Jittered exponential backoff between failed LLM sample retries.

    The free 7B hub is bursty: hammering it after a parse failure tends to
    re-trigger the same 429/500. Wait 1.5s, 3s, 6s, then cap at 12s (with
    +/- 25% jitter) so consecutive samples spread out.
    """
    delay = min(1.5 * (2 ** min(failed_samples, 3)), 12.0)
    jitter = 0.25 * (2 * (time.monotonic() % 1.0) - 1.0)
    time.sleep(max(0.25, delay * (1 + jitter)))


def load_documents(paths: list[str], max_chars: int = 12000) -> str:
    parts: list[str] = []
    for path in paths:
        p = Path(path)
        if not p.exists():
            raise FileNotFoundError(f"Input not found: {path}")
        if p.suffix.lower() in (".pptx", ".ppt"):
            body = _load_pptx(p, max_chars)
        else:
            body = p.read_text(encoding="utf-8")[:max_chars]
        parts.append(f"<doc name='{p.stem}'>\n{body}\n</doc>")
    return "\n\n".join(parts)[: max_chars * 4]
def _load_pptx(p: Path, max_chars: int) -> str:
    from pptx import Presentation

    prs = Presentation(str(p))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        texts = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        if texts:
            parts.append(f"--- Slide {i} ---\n" + "\n".join(texts))
    joined = "\n\n".join(parts)
    if not joined.strip():
        raise ValueError(f"No slide text found in {p}")
    return joined[:max_chars]
